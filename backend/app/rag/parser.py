"""Document parser module supporting PDF, DOCX, TXT, XLSX, and CSV files."""

import os
import re
import csv
from typing import List, Dict, Any
from pypdf import PdfReader
from docx import Document as DocxDocument


class DocumentParser:
    """Extracts raw text, structured tables, and metadata from corporate document files."""

    @staticmethod
    def clean_text(text: str, mask_sensitive_pii: bool = True) -> str:
        """Normalize whitespace, strip control characters, and sanitize sensitive PII."""
        if not text:
            return ""
        # Remove null bytes and non-printable control chars except \n, \t, \r
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]", "", text)
        # Normalize multiple spaces on the same line
        text = re.sub(r"[ \t]+", " ", text)
        # Normalize more than two consecutive newlines into double newlines
        text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
        text = text.strip()

        if mask_sensitive_pii:
            try:
                from app.rag.pii_masker import pii_masker
                text = pii_masker.mask_pii(text)
            except Exception:
                pass

        return text

    @classmethod
    def parse_pdf(cls, file_path: str) -> List[Dict[str, Any]]:
        """Extract text from PDF page by page."""
        pages = []
        reader = PdfReader(file_path)
        for page_idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            cleaned = cls.clean_text(text)
            if cleaned:
                pages.append({
                    "text": cleaned,
                    "page": page_idx + 1,
                })
        return pages

    @classmethod
    def parse_docx(cls, file_path: str) -> List[Dict[str, Any]]:
        """Extract text from DOCX paragraphs and tables formatted as Markdown."""
        doc = DocxDocument(file_path)
        sections = []

        for para in doc.paragraphs:
            text = cls.clean_text(para.text)
            if text:
                sections.append(text)

        for table in doc.tables:
            table_rows = []
            for row_idx, row in enumerate(table.rows):
                cells = [cls.clean_text(cell.text).replace("\n", " ") for cell in row.cells]
                if any(cells):
                    row_str = "| " + " | ".join(cells) + " |"
                    table_rows.append(row_str)
                    if row_idx == 0:
                        separator = "| " + " | ".join(["---"] * len(cells)) + " |"
                        table_rows.append(separator)

            if table_rows:
                sections.append("\n".join(table_rows))

        combined = "\n\n".join(sections)
        cleaned = cls.clean_text(combined)
        return [{"text": cleaned, "page": 1}] if cleaned else []

    @classmethod
    def parse_txt(cls, file_path: str) -> List[Dict[str, Any]]:
        """Read standard UTF-8 text file."""
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
        cleaned = cls.clean_text(content)
        return [{"text": cleaned, "page": 1}] if cleaned else []

    @classmethod
    def parse_csv(cls, file_path: str, rows_per_page: int = 35) -> List[Dict[str, Any]]:
        """Parse CSV file into structured Markdown tables paginated with persistent headers."""
        lines = []
        for encoding in ["utf-8-sig", "utf-8", "cp1252", "latin-1"]:
            try:
                with open(file_path, "r", encoding=encoding) as f:
                    # Detect delimiter
                    sample = f.read(4096)
                    f.seek(0)
                    delimiter = ","
                    try:
                        dialect = csv.Sniffer().sniff(sample)
                        delimiter = dialect.delimiter
                    except Exception:
                        pass
                    reader = csv.reader(f, delimiter=delimiter)
                    lines = [[cls.clean_text(cell) for cell in row] for row in reader if any(row)]
                if lines:
                    break
            except UnicodeDecodeError:
                continue

        if not lines:
            return []

        header = lines[0]
        data_rows = lines[1:]

        if not data_rows:
            header_str = "| " + " | ".join(header) + " |\n| " + " | ".join(["---"] * len(header)) + " |"
            return [{"text": header_str, "page": 1}]

        pages = []
        page_num = 1
        for i in range(0, len(data_rows), rows_per_page):
            chunk_rows = data_rows[i:i + rows_per_page]
            table_lines = [
                "| " + " | ".join(header) + " |",
                "| " + " | ".join(["---"] * len(header)) + " |"
            ]
            for r in chunk_rows:
                # Pad or trim row to match header length
                padded = r + [""] * (len(header) - len(r)) if len(r) < len(header) else r[:len(header)]
                table_lines.append("| " + " | ".join(padded) + " |")

            page_text = "\n".join(table_lines)
            pages.append({
                "text": page_text,
                "page": page_num,
            })
            page_num += 1

        return pages

    @classmethod
    def parse_xlsx(cls, file_path: str, rows_per_page: int = 35) -> List[Dict[str, Any]]:
        """Parse Excel workbook (.xlsx) sheets into Markdown tables paginated with headers."""
        import openpyxl

        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        pages = []
        page_num = 1

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            all_rows = []
            for row in sheet.iter_rows(values_only=True):
                if row and any(v is not None and str(v).strip() for v in row):
                    cleaned_row = [cls.clean_text(str(v or "")).replace("\n", " ") for v in row]
                    all_rows.append(cleaned_row)

            if not all_rows:
                continue

            header = all_rows[0]
            data_rows = all_rows[1:]

            if not data_rows:
                header_str = f"### Trang tính: {sheet_name}\n\n| " + " | ".join(header) + " |\n| " + " | ".join(["---"] * len(header)) + " |"
                pages.append({"text": header_str, "page": page_num})
                page_num += 1
                continue

            for i in range(0, len(data_rows), rows_per_page):
                chunk_rows = data_rows[i:i + rows_per_page]
                table_lines = [
                    f"### Trang tính: {sheet_name}",
                    "| " + " | ".join(header) + " |",
                    "| " + " | ".join(["---"] * len(header)) + " |"
                ]
                for r in chunk_rows:
                    padded = r + [""] * (len(header) - len(r)) if len(r) < len(header) else r[:len(header)]
                    table_lines.append("| " + " | ".join(padded) + " |")

                page_text = "\n".join(table_lines)
                pages.append({
                    "text": page_text,
                    "page": page_num,
                })
                page_num += 1

        wb.close()
        return pages

    @classmethod
    def parse_file(cls, file_path: str, file_type: str) -> List[Dict[str, Any]]:
        """Dispatch parsing based on file extension."""
        ft = file_type.upper()
        if ft == "PDF":
            return cls.parse_pdf(file_path)
        elif ft in ["DOCX", "DOC"]:
            return cls.parse_docx(file_path)
        elif ft == "TXT":
            return cls.parse_txt(file_path)
        elif ft in ["XLSX", "XLS"]:
            return cls.parse_xlsx(file_path)
        elif ft == "CSV":
            return cls.parse_csv(file_path)
        else:
            raise ValueError(f"Định dạng tệp '{file_type}' không được hỗ trợ. Chỉ hỗ trợ PDF, DOCX, TXT, XLSX, CSV.")


parser = DocumentParser()
